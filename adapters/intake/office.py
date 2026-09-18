# -*- coding: utf-8 -*-
"""OOXML formats: DOCX, PPTX, XLSX.

All three are ZIP packages, which is why identifying them needs more than magic bytes. A
``.docx`` renamed to ``.xlsx`` has identical leading bytes; only the package contents say what
it really is. :func:`core.intake.policy.resolve_ooxml` holds that rule and this module applies
it, on the package it has actually opened, before handing anything to a parsing library.

The same single pass over the central directory also enforces the uncompressed-size limit. A
ZIP declares each entry's expanded size there, so the total is known before a byte is
decompressed.
"""
from __future__ import annotations

import zipfile
from io import BytesIO

from adapters.intake.guard import parser_guard
from core.errors import IntakeError, IntakeErrorCode
from core.intake.extract import make_segment
from core.intake.models import DocumentSegment, ExtractedDocument, SegmentKind
from core.intake.policy import (
    DEFAULT_POLICY,
    IntakePolicy,
    check_archive_size,
    resolve_ooxml,
)
from core.models import FileType

#: Rows per spreadsheet segment.
XLSX_ROWS_PER_SEGMENT = 50


def verify_ooxml_package(data: bytes, declared: FileType, policy: IntakePolicy) -> None:
    """Confirm the package really is the declared type, and is not a decompression bomb.

    Raises ``MALFORMED_DOCUMENT``, ``ARCHIVE_LIMIT_EXCEEDED`` or ``TYPE_MISMATCH``.
    """
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            infos = archive.infolist()
    except zipfile.BadZipFile:
        raise IntakeError(IntakeErrorCode.MALFORMED_DOCUMENT) from None
    except Exception:  # noqa: BLE001
        raise IntakeError(IntakeErrorCode.MALFORMED_DOCUMENT) from None

    # Cheap, and it runs before any entry is expanded.
    check_archive_size((info.file_size for info in infos), policy)

    actual = resolve_ooxml(info.filename for info in infos)
    if actual is None or actual is not declared:
        raise IntakeError(IntakeErrorCode.TYPE_MISMATCH)


def _docx_segments(data: bytes) -> tuple[list[DocumentSegment], int | None]:
    try:
        from docx import Document
    except ImportError:
        raise IntakeError(IntakeErrorCode.PARSER_UNAVAILABLE, parser="office") from None

    segments: list[DocumentSegment] = []
    order = 0

    with parser_guard("office"):
        document = Document(BytesIO(data))

        for index, paragraph in enumerate(document.paragraphs, start=1):
            style = (paragraph.style.name or "") if paragraph.style is not None else ""
            kind = SegmentKind.HEADING if style.startswith("Heading") else SegmentKind.PARAGRAPH
            segment = make_segment(paragraph.text, f"¶{index}", kind=kind, order=order)
            if segment is not None:
                segments.append(segment)
                order += 1

        for index, table in enumerate(document.tables, start=1):
            rows = [
                " | ".join(cell.text.strip() for cell in row.cells) for row in table.rows
            ]
            segment = make_segment(
                "\n".join(rows), f"table {index}", kind=SegmentKind.TABLE, order=order
            )
            if segment is not None:
                segments.append(segment)
                order += 1

    # A Word document has no stable page count without rendering it, so none is reported
    # rather than a number that would not survive a different renderer.
    return segments, None


def _pptx_segments(data: bytes) -> tuple[list[DocumentSegment], int | None]:
    try:
        from pptx import Presentation
    except ImportError:
        raise IntakeError(IntakeErrorCode.PARSER_UNAVAILABLE, parser="office") from None

    segments: list[DocumentSegment] = []
    order = 0

    with parser_guard("office"):
        presentation = Presentation(BytesIO(data))
        slide_count = len(presentation.slides)

        for number, slide in enumerate(presentation.slides, start=1):
            body: list[str] = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    text = shape.text_frame.text
                    if text.strip():
                        body.append(text)

            segment = make_segment(
                "\n".join(body), f"slide {number}", kind=SegmentKind.PARAGRAPH, order=order
            )
            if segment is not None:
                segments.append(segment)
                order += 1

            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text
                note_segment = make_segment(
                    notes, f"slide {number} notes", kind=SegmentKind.SLIDE_NOTE, order=order
                )
                if note_segment is not None:
                    segments.append(note_segment)
                    order += 1

    return segments, slide_count


def _xlsx_segments(data: bytes) -> tuple[list[DocumentSegment], int | None]:
    try:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise IntakeError(IntakeErrorCode.PARSER_UNAVAILABLE, parser="office") from None

    segments: list[DocumentSegment] = []
    order = 0

    with parser_guard("office"):
        # read_only keeps memory bounded on large sheets; data_only takes cached formula
        # results rather than the formula text, which is what an analyst actually reads.
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
        try:
            sheet_count = len(workbook.worksheets)
            for sheet in workbook.worksheets:
                rows = [
                    [("" if value is None else str(value)).strip() for value in row]
                    for row in sheet.iter_rows(values_only=True)
                ]
                rows = [row for row in rows if any(cell for cell in row)]
                if not rows:
                    continue

                width = max(len(row) for row in rows)
                last_column = get_column_letter(max(width, 1))

                for offset in range(0, len(rows), XLSX_ROWS_PER_SEGMENT):
                    chunk = rows[offset : offset + XLSX_ROWS_PER_SEGMENT]
                    start = offset + 1
                    end = start + len(chunk) - 1
                    locator = f"{sheet.title}!A{start}:{last_column}{end}"
                    body = "\n".join(" | ".join(row) for row in chunk)
                    segment = make_segment(
                        body, locator, kind=SegmentKind.CELL_RANGE, order=order
                    )
                    if segment is not None:
                        segments.append(segment)
                        order += 1
        finally:
            workbook.close()

    return segments, sheet_count


class OfficeParser:
    """DOCX, PPTX and XLSX. Implements :class:`core.interfaces.intake.DocumentParser`."""

    name = "office"
    supported_types = frozenset({FileType.DOCX, FileType.PPTX, FileType.XLSX})

    def __init__(self, policy: IntakePolicy = DEFAULT_POLICY) -> None:
        self.policy = policy

    def parse(self, data: bytes, *, file_type: FileType) -> ExtractedDocument:
        if file_type not in self.supported_types:
            raise IntakeError(IntakeErrorCode.UNSUPPORTED_FILE_TYPE, parser=self.name)

        verify_ooxml_package(data, file_type, self.policy)

        if file_type is FileType.DOCX:
            segments, page_count = _docx_segments(data)
        elif file_type is FileType.PPTX:
            segments, page_count = _pptx_segments(data)
        else:
            segments, page_count = _xlsx_segments(data)

        return ExtractedDocument(
            file_type=file_type,
            byte_size=len(data),
            segments=segments,
            page_count=page_count,
            detected_lang=None,
            parser_name=self.name,
        )
